"""Export every AccuLynx job location as one row per property (address + unit), in the
parcel-enrichment template's column order (Address, Unit#, City, State, Zip, County, FIPS,
APN#) plus a Ref ID, with a job crosswalk sheet. See docs/113-property-spine.md.

Usage:  python3 scripts/export-acculynx-property-addresses.py [--out path.xlsx]

Reads SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY from the environment, else from .env at the
repo root, else from the main checkout's .env (worktrees). Values are never printed.
PostgREST truncates silently at 1,000 rows, so every read paginates (CONVENTIONS §10).
Read-only: this script never writes to the database.
"""
import json, os, re, subprocess, sys, urllib.request, urllib.parse
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[sys.argv.index("--out") + 1]) if "--out" in sys.argv else ROOT / "exports" / "acculynx-job-properties.xlsx"

def env_files():
    yield ROOT / ".env"
    try:
        common = subprocess.run(["git", "rev-parse", "--git-common-dir"], cwd=ROOT, capture_output=True, text=True).stdout.strip()
        if common:
            yield (ROOT / common).resolve().parent / ".env"
    except OSError:
        pass

env = {}
for f in env_files():
    if f.is_file():
        for line in f.read_text().splitlines():
            m = re.match(r"^\s*([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.*)\s*$", line)
            if m:
                env.setdefault(m[1], m[2].strip().strip("'\""))
        break
URL = os.environ.get("SUPABASE_URL") or env.get("SUPABASE_URL")
KEY = os.environ.get("SUPABASE_SERVICE_ROLE_KEY") or env.get("SUPABASE_SERVICE_ROLE_KEY")
if not URL or not KEY:
    sys.exit("missing SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY")

SELECT = ",".join([
    "id", "job_number", "job_name", "job_category_name", "current_milestone", "created_date",
    "location_street1", "location_city", "location_state_abbrev", "location_zip",
    "latitude", "longitude",
    "street2:raw->locationAddress->>street2",
    "state_name:raw->locationAddress->state->>name",
])

def fetch_all():
    rows, page = [], 1000
    for off in range(0, 10**6, page):
        q = urllib.parse.urlencode({"select": SELECT, "order": "id.asc", "limit": page, "offset": off})
        req = urllib.request.Request(f"{URL}/rest/v1/acculynx_jobs?{q}",
                                     headers={"apikey": KEY, "Authorization": f"Bearer {KEY}"})
        with urllib.request.urlopen(req, timeout=60) as r:
            batch = json.load(r)
        rows += batch
        if len(batch) < page:
            return rows

ws_re = re.compile(r"\s+")
UNIT_RE = re.compile(
    r"(?:\s*,\s*|\s+)(?P<unit>(?:#\s*|(?:apt|apartment|unit|ste|suite|bldg|building|lot|spc|space|fl|floor|rm|room)\.?\s*#?\s*)[\w-]+)\s*$",
    re.I)

def clean(s):
    return ws_re.sub(" ", (s or "").strip())

def split_address(s1, s2):
    s1, s2 = clean(s1), clean(s2)
    notes = []
    # street2 holding "ST 12345" (state + zip) or a PO box is not a unit — drop it
    if s2 and (re.fullmatch(r"[A-Za-z]{2}\.?\s+\d{5}(-\d{4})?", s2) or re.match(r"^p\.?\s*o\.?\s*box", s2, re.I)):
        notes.append(f"dropped street2 '{s2}'")
        s2 = ""
    # House number typed into street2 (street1 "Roan Place", street2 "16423")
    if s2 and re.fullmatch(r"\d+[A-Za-z]?(-\d+)?", s2) and not re.match(r"^\d", s1):
        return f"{s2} {s1}".strip(), "", "; ".join(notes + ["number from street2"])
    # Street name typed into street2 (street1 "29113", street2 "Amerind Springs Trail")
    if s2 and re.fullmatch(r"\d+[A-Za-z]?", s1) and re.match(r"^[A-Za-z]", s2):
        return f"{s1} {s2}", "", "; ".join(notes + ["street name from street2"])
    unit, note = "", "; ".join(notes)
    m = UNIT_RE.search(s1)
    if m and re.match(r"^\d", s1):
        unit = m["unit"].lstrip("#").strip()
        s1 = s1[: m.start()].rstrip(" ,")
        note = "unit parsed from street1"
    if s2:
        unit = f"{unit} {s2.lstrip('#')}".strip() if unit else s2.lstrip("#").strip()
        note = (note + "; " if note else "") + "unit from street2"
    return s1, unit, note

def zip5(z):
    d = re.sub(r"[^0-9]", "", z or "")
    if not d:
        return ""
    return d[:5].zfill(5) if len(d) <= 5 else d[:5]

TAIL_RE = re.compile(r"[\s,]+(?P<st>[A-Za-z]{2})\.?\s+(?P<zip>\d{5})(?:-\d{4})?\s*$")

jobs = fetch_all()
for j in jobs:
    # street1 carrying its own "..., City ST 12345" tail: lift state/zip out before unit parsing
    s1 = clean(j.get("location_street1"))
    m = TAIL_RE.search(s1)
    if m and re.match(r"^\d", s1):
        j["location_street1"] = s1[: m.start()]
        j["location_state_abbrev"] = j.get("location_state_abbrev") or m["st"].upper()
        j["location_zip"] = j.get("location_zip") or m["zip"]
        j["_tail_note"] = "state/zip from street1"
print(f"fetched {len(jobs)} jobs")

props = {}          # key -> property dict
job_rows = []
missing = []
for j in jobs:
    addr, unit, note = split_address(j.get("location_street1"), j.get("street2"))
    city = clean(j.get("location_city"))
    # "1451 Middle Gulf Dr, Sanibel" with no city: the city rode along in street1
    if not city and "," in addr.rstrip(" ,"):
        addr, city = [clean(x) for x in addr.rstrip(" ,").rsplit(",", 1)]
        note = (note + "; " if note else "") + "city from street1"
    addr = addr.rstrip(" ,")
    if j.get("_tail_note"):
        note = (note + "; " if note else "") + j["_tail_note"]
    state = clean(j.get("location_state_abbrev")) or clean(j.get("state_name"))
    z = zip5(j.get("location_zip"))
    base = {
        "job": j, "addr": addr, "unit": unit, "city": city, "state": state, "zip": z, "note": note,
    }
    if not addr:
        missing.append(base)
        continue
    key = (addr.lower(), unit.lower(), z or city.lower(), state.upper())
    p = props.setdefault(key, {"addr": addr, "unit": unit, "city": city, "state": state, "zip": z,
                               "jobs": [], "cats": Counter()})
    p["jobs"].append(j)
    if j.get("job_category_name"):
        p["cats"][j["job_category_name"]] += 1
    job_rows.append((key, base))

# Stable Ref IDs grouped by market: state, city, address
ordered = sorted(props.items(), key=lambda kv: (kv[1]["state"], kv[1]["city"].lower(), kv[1]["addr"].lower(), kv[1]["unit"]))
ref_of = {}
for i, (k, p) in enumerate(ordered, 1):
    ref_of[k] = f"PX-{i:05d}"

HEAD_FILL = PatternFill("solid", fgColor="11133F")
HEAD_FONT = Font(bold=True, color="FFFFFF")

def sheet(ws, headers, rows, widths, text_cols=()):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col in text_cols:  # keep ZIP / FIPS / APN as text so leading zeros survive
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for cell in row:
                cell.number_format = "@"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

wb = Workbook()
ws = wb.active
ws.title = "properties"
sheet(ws,
      ["Address", "Unit#", "City", "State", "Zip", "County", "FIPS", "APN#", "Ref ID"],
      [[p["addr"], p["unit"], p["city"], p["state"], p["zip"], "", "", "", ref_of[k]] for k, p in ordered],
      [38, 10, 20, 7, 8, 18, 8, 22, 11], text_cols=(5, 7, 8))

ws2 = wb.create_sheet("jobs")
jr = []
for key, b in sorted(job_rows, key=lambda kb: (ref_of[kb[0]], kb[1]["job"].get("job_number") or "")):
    j = b["job"]
    jr.append([ref_of[key], j["id"], j.get("job_number"), j.get("job_name"), j.get("job_category_name"),
               j.get("current_milestone"), (j.get("created_date") or "")[:10],
               j.get("latitude"), j.get("longitude"),
               clean(j.get("location_street1")), clean(j.get("street2")), b["note"]])
sheet(ws2, ["Ref ID", "AccuLynx Job ID", "Job Number", "Job Name", "Job Category", "Milestone", "Created",
            "Latitude", "Longitude", "AccuLynx Street 1", "AccuLynx Street 2", "Cleanup Note"],
      jr, [11, 38, 14, 34, 18, 16, 11, 11, 12, 34, 16, 26])

ws3 = wb.create_sheet("needs_address")
sheet(ws3, ["AccuLynx Job ID", "Job Number", "Job Name", "Job Category", "City", "State", "Zip", "Latitude", "Longitude"],
      [[b["job"]["id"], b["job"].get("job_number"), b["job"].get("job_name"), b["job"].get("job_category_name"),
        b["city"], b["state"], b["zip"], b["job"].get("latitude"), b["job"].get("longitude")] for b in missing],
      [38, 14, 34, 18, 18, 7, 8, 11, 12], text_cols=(7,))

by_state = defaultdict(lambda: [0, 0])
for k, p in ordered:
    by_state[p["state"] or "(blank)"][0] += 1
    by_state[p["state"] or "(blank)"][1] += len(p["jobs"])
ws4 = wb.create_sheet("summary")
sheet(ws4, ["State", "Properties", "Jobs"],
      sorted(([s, v[0], v[1]] for s, v in by_state.items()), key=lambda r: -r[1]) +
      [["Total", len(ordered), len(job_rows)], ["Jobs without a street address (needs_address)", None, len(missing)]],
      [44, 12, 10])

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(json.dumps({
    "jobs": len(jobs), "properties": len(ordered), "jobs_on_properties": len(job_rows), "needs_address": len(missing),
    "multi_job_properties": sum(1 for _, p in ordered if len(p["jobs"]) > 1),
    "with_unit": sum(1 for _, p in ordered if p["unit"]),
    "number_from_street2": sum(1 for _, b in job_rows if b["note"].startswith("number from street2")),
    "no_zip": sum(1 for _, p in ordered if not p["zip"]),
    "no_state": sum(1 for _, p in ordered if not p["state"]),
    "states": {s: v[0] for s, v in sorted(by_state.items(), key=lambda kv: -kv[1][0])},
    "out": str(OUT),
}, indent=1))
