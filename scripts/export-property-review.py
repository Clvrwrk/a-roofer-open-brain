"""Export the AccuLynx job→property review queue to a workbook (docs/115).

Usage:  python3 scripts/export-property-review.py [--out path.xlsx]

Sheets:
  summary               counts per recommended action and what each one asks for
  resubmit_enrichment   real, precisely geocoded addresses the enrichment vendor did not return,
                        in the enrichment template (Address, Unit#, City, State, Zip, County,
                        FIPS, APN#) using Google's standardized address, plus Ref ID
  resubmit_jobs         Ref ID → AccuLynx jobs, to match the vendor's answer back
  human_review          every other open row, with blank decision columns
  acculynx_zip_fixes    jobs linked through the geocoder whose AccuLynx zip is wrong
  fuzzy_links_check     jobs auto-linked by street-name similarity, for spot-checking

Read-only against the database. Keys from env / repo .env / the main checkout's .env.
"""
import json, os, re, subprocess, sys, urllib.request
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[sys.argv.index("--out") + 1]) if "--out" in sys.argv else \
    ROOT / "exports" / f"acculynx-property-review-{date.today().isoformat()}.xlsx"

def load_env():
    env = {}
    files = [ROOT / ".env"]
    try:
        common = subprocess.run(["git", "rev-parse", "--git-common-dir"], cwd=ROOT, capture_output=True, text=True).stdout.strip()
        if common:
            files.append((ROOT / common).resolve().parent / ".env")
    except OSError:
        pass
    for f in files:
        if f.is_file():
            for line in f.read_text().splitlines():
                m = re.match(r"^\s*([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.*)\s*$", line)
                if m:
                    env.setdefault(m[1], m[2].strip().strip("'\""))
            break
    url = os.environ.get("SUPABASE_URL") or env.get("SUPABASE_URL")
    key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY") or env.get("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        sys.exit("missing SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY")
    return url, key

URL, KEY = load_env()

def all_rows(path):
    out = []
    for off in range(0, 10**7, 1000):  # PostgREST truncates silently at 1,000 — always page
        sep = "&" if "?" in path else "?"
        req = urllib.request.Request(f"{URL}/rest/v1/{path}{sep}limit=1000&offset={off}",
                                     headers={"apikey": KEY, "Authorization": f"Bearer {KEY}"})
        with urllib.request.urlopen(req, timeout=120) as r:
            page = json.load(r)
        out += page
        if len(page) < 1000:
            return out

def county_key(name):  # mirrors public.county_key() (mig 304)
    s = re.sub(r"[^A-Za-z0-9 ]", "", name or "").upper()
    s = re.sub(r"(^|\s)SAINTE?(\s)", r"\1ST\2", s)
    s = re.sub(r"\s+(COUNTY|PARISH|BOROUGH|CENSUS AREA|MUNICIPALITY)$", "", s)
    return re.sub(r"\s+", " ", s).strip()

ACTIONS = {
    "resubmit_for_enrichment": "Real address (Google pins it to the rooftop or street range) that the enrichment vendor did not return. Send the resubmit_enrichment sheet back to the vendor; the answer loads with the same scripts.",
    "fix_address_in_acculynx": "The address in AccuLynx is missing, incomplete, or cannot be located. Correct it on the AccuLynx job; the next link run picks it up automatically.",
    "confirm_suggested_match": "A likely property exists but the street name differs too much to link automatically. Confirm the suggested property or give the right one.",
    "choose_property": "Several properties share this street address (units or condos). Pick the one the job is for.",
    "confirm_not_a_property": "The job does not describe a real property (test or placeholder). Confirm so it can be dismissed.",
}

def main():
    review = all_rows("acculynx_job_property_review?select=*&order=job_id")
    jobs = {j["id"]: j for j in all_rows("acculynx_jobs?select=id,job_number,job_name,job_category_name,current_milestone,created_date,location_street1,location_zip,property_id,property_link_method&order=id")}
    props = {p["id"]: p for p in all_rows("properties?select=id,address_full,address_key&order=id")}
    geo = {g["target_id"]: g for g in all_rows("geocode_result?select=target_id,county_name,postal_code,formatted_address&target_kind=eq.acculynx_job")}
    fips = {(c["state_abbrev"], c["county_key"]): c["county_fips"] for c in all_rows("county_ref?select=state_abbrev,county_key,county_fips")}
    open_rows = [r for r in review if r["status"] == "open"]

    HEAD = PatternFill("solid", fgColor="11133F")
    def sheet(ws, headers, rows, widths, text_cols=()):
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill, cell.font = HEAD, Font(bold=True, color="FFFFFF")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for col in text_cols:
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                for cell in row:
                    cell.number_format = "@"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    def job_link(jid):
        return f"https://my.acculynx.com/jobs/{jid}"

    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    counts = Counter(r["recommended_action"] for r in open_rows)
    sheet(ws, ["Recommended action", "Jobs", "What it asks for"],
          [[a, counts.get(a, 0), d] for a, d in ACTIONS.items()] +
          [["Total open", len(open_rows), ""],
           ["Linked so far", sum(1 for j in jobs.values() if j["property_id"]), f"of {len(jobs)} AccuLynx jobs"]],
          [28, 8, 110])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Resubmit: one row per standardized address, in the enrichment template.
    groups = defaultdict(list)
    for r in open_rows:
        if r["recommended_action"] != "resubmit_for_enrichment" or not r.get("geocoded_address"):
            continue
        parts = [p.strip() for p in r["geocoded_address"].split(",")]
        if len(parts) < 3:
            continue
        m = re.match(r"^([A-Z]{2})\s+(\d{5})", parts[-2] if parts[-1] == "USA" else parts[-1])
        street, city = parts[0], parts[1] if len(parts) > 3 else ""
        st, zp = (m.group(1), m.group(2)) if m else ("", "")
        unit = ""
        um = re.search(r"\s*(?:#|Unit|Ste|Suite|Apt)\s*([\w-]+)$", street)
        if um:
            unit, street = um.group(1), street[: um.start()].strip()
        g = geo.get(r["job_id"], {})
        county = re.sub(r"\s+(County|Parish)$", "", g.get("county_name") or "")
        groups[(street.upper(), unit.upper(), zp)].append((r, street, unit, city, st, zp, county))
    resub, resub_jobs = [], []
    for i, (k, items) in enumerate(sorted(groups.items(), key=lambda kv: (kv[1][0][4], kv[1][0][3], kv[0][0])), 1):
        r, street, unit, city, st, zp, county = items[0]
        ref = f"RV-{i:05d}"
        resub.append([street, unit, city, st, zp, county, fips.get((st, county_key(county)), ""), "", ref])
        for (rr, *_rest) in items:
            j = jobs.get(rr["job_id"], {})
            resub_jobs.append([ref, rr["job_id"], j.get("job_number"), j.get("job_name"), rr["submitted_address"],
                               rr["geocoded_address"]])
    ws = wb.create_sheet("resubmit_enrichment")
    sheet(ws, ["Address", "Unit#", "City", "State", "Zip", "County", "FIPS", "APN#", "Ref ID"], resub,
          [34, 8, 18, 6, 8, 16, 8, 16, 10], text_cols=(5, 7, 8))
    ws = wb.create_sheet("resubmit_jobs")
    sheet(ws, ["Ref ID", "AccuLynx Job ID", "Job Number", "Job Name", "AccuLynx Address", "Google Standardized Address"],
          resub_jobs, [10, 38, 12, 34, 44, 44])

    # Human review: everything else, with decision columns.
    order = {a: i for i, a in enumerate(ACTIONS)}
    hr = []
    for r in sorted((r for r in open_rows if r["recommended_action"] != "resubmit_for_enrichment"),
                    key=lambda r: (order.get(r["recommended_action"], 9), r["reason"], r["submitted_address"] or "")):
        j = jobs.get(r["job_id"], {})
        sp = props.get(r.get("suggested_property_id") or "", {})
        hr.append([r["recommended_action"], r["reason"], job_link(r["job_id"]), j.get("job_number"), j.get("job_name"),
                   j.get("job_category_name"), j.get("current_milestone"), (j.get("created_date") or "")[:10],
                   r["submitted_address"], r.get("geocoded_address"), r.get("geocode_precision"),
                   sp.get("address_full"), r.get("suggestion_score"), "", "", ""])
    ws = wb.create_sheet("human_review")
    sheet(ws, ["Action", "Reason", "AccuLynx Job", "Job Number", "Job Name", "Category", "Milestone", "Created",
               "AccuLynx Address", "Google Standardized Address", "Google Precision", "Suggested Property",
               "Match Score", "Decision (link / fix / not a property)", "Corrected Address or Property", "Notes"],
          hr, [24, 22, 30, 11, 30, 14, 14, 10, 40, 40, 16, 36, 8, 22, 36, 30])

    zf = []
    for r in review:
        res = r.get("resolution") or ""
        m = re.search(r"AccuLynx zip (\S+) should be (\d{5})", res)
        if r["status"] == "resolved" and m:
            j = jobs.get(r["job_id"], {})
            zf.append([job_link(r["job_id"]), j.get("job_number"), j.get("job_name"), r["submitted_address"],
                       m.group(1), m.group(2), (props.get(r.get("resolved_property_id") or "", {}) or {}).get("address_full")])
    ws = wb.create_sheet("acculynx_zip_fixes")
    sheet(ws, ["AccuLynx Job", "Job Number", "Job Name", "AccuLynx Address", "AccuLynx Zip", "Correct Zip", "Linked Property"],
          zf, [30, 11, 30, 44, 11, 11, 40], text_cols=(5, 6))

    fz = []
    for j in jobs.values():
        if j.get("property_link_method") == "address_fuzzy":
            fz.append([job_link(j["id"]), j.get("job_number"), j.get("job_name"),
                       f"{j.get('location_street1') or ''} {j.get('location_zip') or ''}".strip(),
                       (props.get(j["property_id"], {}) or {}).get("address_full"), "", ""])
    ws = wb.create_sheet("fuzzy_links_check")
    sheet(ws, ["AccuLynx Job", "Job Number", "Job Name", "AccuLynx Address", "Linked Property", "Wrong? (Y)", "Notes"],
          sorted(fz, key=lambda r: r[3]), [30, 11, 30, 40, 40, 10, 30])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(json.dumps({"out": str(OUT), "open": len(open_rows), "by_action": counts,
                      "resubmit_addresses": len(resub), "resubmit_jobs": len(resub_jobs), "human_review": len(hr),
                      "zip_fixes": len(zf), "fuzzy_links": len(fz)}, default=str))

if __name__ == "__main__":
    main()
