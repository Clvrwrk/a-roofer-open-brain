"""Stage a property-list export (the list tool's 75-column xlsx) into property_list_import.

Usage:
  python3 scripts/stage-property-list.py <import_batch> <file.xlsx>[=<market>] [<file2.xlsx>[=<market>] ...]

Each row is stored verbatim as `raw` jsonb, keyed (import_batch, source_file, row_number), so
re-staging the same file is an idempotent upsert. Then load it (docs/114 §6, docs/116):
  SELECT public.load_property_list_import('<import_batch>');

Reads SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY from env, repo .env, or the main checkout's
.env; values are never printed. This script only writes the staging table.
"""
import datetime, json, os, re, subprocess, sys, urllib.request, warnings
from pathlib import Path
from openpyxl import load_workbook

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parent.parent

def load_env():
    env = {}
    files = [ROOT / ".env"]
    try:
        common = subprocess.run(["git", "rev-parse", "--git-common-dir"], cwd=ROOT, capture_output=True, text=True).stdout.strip()
        if common:
            files.append((ROOT / common).resolve().parent / ".env")
    except OSError:
        pass
    for f in files:
        if f.is_file():
            for line in f.read_text().splitlines():
                m = re.match(r"^\s*([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.*)\s*$", line)
                if m:
                    env.setdefault(m[1], m[2].strip().strip("'\""))
            break
    url = os.environ.get("SUPABASE_URL") or env.get("SUPABASE_URL")
    key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY") or env.get("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        sys.exit("missing SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY")
    return url, key

def val(v):
    return v.isoformat() if isinstance(v, (datetime.datetime, datetime.date)) else v

def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    url, key = load_env()
    batch = sys.argv[1]
    out = []
    for spec in sys.argv[2:]:
        path, _, market = spec.partition("=")
        p = Path(path)
        ws = load_workbook(p, data_only=True).worksheets[0]  # full load: some exports lack dimension metadata
        it = ws.iter_rows(values_only=True)
        hdr = next(it)
        for i, r in enumerate(it, start=2):
            raw = {h: val(v) for h, v in zip(hdr, r) if h is not None}
            if not any(str(v or "").strip() for v in raw.values()):
                continue
            out.append({"import_batch": batch, "source_file": p.name, "row_number": i, "market": market or None,
                        "apn": (str(raw.get("APN") or "").strip() or None), "county": raw.get("County") or None,
                        "raw": raw})
    for k in range(0, len(out), 500):
        req = urllib.request.Request(
            f"{url}/rest/v1/property_list_import?on_conflict=import_batch,source_file,row_number",
            data=json.dumps(out[k:k + 500]).encode(), method="POST",
            headers={"apikey": key, "Authorization": f"Bearer {key}", "Content-Type": "application/json",
                     "Prefer": "resolution=merge-duplicates,return=minimal"})
        with urllib.request.urlopen(req, timeout=120) as resp:
            assert resp.status in (200, 201, 204), resp.status
    print(json.dumps({"batch": batch, "rows_staged": len(out)}))

if __name__ == "__main__":
    main()
